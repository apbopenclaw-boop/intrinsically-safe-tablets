#!/usr/bin/env python3
"""Import the tablet catalogue from the master data workbook (explosionproof_tablets_master_data.xlsx).

    npm run tablets:import -- /path/to/explosionproof_tablets_master_data.xlsx
    npm run tablets:import -- /path/to/workbook.xlsx --report notes-report.md   # list every public note, don't write

Needs openpyxl (python3 -m pip install openpyxl). The workbook is not kept in this public repo: it holds review notes
and prices that must not be published.

Writes one JSON file per published tablet to src/content/tablets/ and deletes files for tablets that are no longer
published. Images (image, imageAlt) added by hand to an existing file are kept; everything else comes from the
workbook, so change the workbook or the rules below rather than the JSON.

Publishing rules (same as the phones site, plus the tablet rules agreed in ../research/spec/verify-4.json)
- Announced products are left out until they are certified (IS941.1, Atexxo iPad 11" Zone 1). Rows with Page depth
  "Variant of <ID>" are listed as versions on their parent page (mining versions, the D2 North America version, the
  Mobile Operator Panel, Holland Shielding's listing of the Atexxo iPad mini 5).
- NOT FOUND cells are never published. CHECK and CONFLICT values are published with a public note from NOTE_RULES or
  FIELD_NOTES. The import stops if a published flagged cell has no note, so researcher comments never reach the site.
- Some flagged cells are never published (SKIP_RULES): marketing "equivalent to Class I Div 2" claims, UKCA mentions
  without a certificate, and "no successor" placeholders.
- Never call a product intrinsically safe unless its marking is Ex i. Flameproof (Ex db), restricted-breathing (Ex nR)
  and Ex ec products are described by their concept.
- Summaries, strengths, limitations, alternatives and best-for tags follow the same rules for every brand.
"""

import argparse
import datetime
import json
import re
import sys
from collections import Counter, namedtuple
from fnmatch import fnmatch
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: python3 -m pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "src" / "content" / "tablets"

Cell = namedtuple("Cell", "v flag comment")
FLAG_FILLS = {"FFF2CC": "check", "F8CBAD": "conflict", "E7E6E6": "notfound"}
PLACEHOLDER = re.compile(
    r"^(none\b.*|n/a\b.*|not (found|stated|listed|verified|published|confirmed|applicable)\b.*|check|unknown|-|—)$", re.I
)
HEDGED = re.compile(r"not verified|not stated|not confirmed|assumed|likely|implied|possible|inferred", re.I)

# ---------------------------------------------------------------------------------------------------------------
# Catalogue: which workbook rows become pages, and under which URL and name
# ---------------------------------------------------------------------------------------------------------------

PRODUCTS = {
    # i.safe MOBILE
    "IS-930.1": ("isafe-is930-1", "i.safe MOBILE IS930.1"),
    "IS-930.2": ("isafe-is930-2", "i.safe MOBILE IS930.2"),
    "IS-940.1": ("isafe-is940-1", "i.safe MOBILE IS940.1"),
    "IS-940.2": ("isafe-is940-2", "i.safe MOBILE IS940.2"),
    "IS-945.1": ("isafe-is945-1", "i.safe MOBILE IS945.1"),
    "IS-945.2": ("isafe-is945-2", "i.safe MOBILE IS945.2"),
    "IS-910.1": ("isafe-is910-1", "i.safe MOBILE IS910.1"),
    "IS-910.2": ("isafe-is910-2", "i.safe MOBILE IS910.2"),
    # ecom / Pepperl+Fuchs
    "PF-TE05-DZ1": ("ecom-tab-ex-05-dz1", "ecom Tab-Ex 05 DZ1"),
    "PF-TE05-DZ2": ("ecom-tab-ex-05-dz2", "ecom Tab-Ex 05 DZ2"),
    "PF-PADEX01-P12-DZ2": ("ecom-pad-ex-01-p12-dz2", "ecom Pad-Ex 01 P12 DZ2"),
    "PF-TE02-DZ1": ("ecom-tab-ex-02-dz1", "ecom Tab-Ex 02 DZ1"),
    "PF-TE02-DZ2": ("ecom-tab-ex-02-dz2", "ecom Tab-Ex 02 DZ2"),
    "PF-TE03-DZ1": ("ecom-tab-ex-03-dz1", "ecom Tab-Ex 03 DZ1"),
    "PF-TE03-DZ2": ("ecom-tab-ex-03-dz2", "ecom Tab-Ex 03 DZ2"),
    "PF-TEPRO-DZ2": ("ecom-tab-ex-pro-dz2", "ecom Tab-Ex Pro DZ2"),
    "PF-TE04PRO-DZ2": ("ecom-tab-ex-04-pro-dz2", "ecom Tab-Ex 04 Pro DZ2"),
    "PF-PADEX01-P8": ("ecom-pad-ex-01", "ecom Pad-Ex 01"),
    # BARTEC
    "BT-AGILE-S-NI": ("bartec-agile-s-ni", "BARTEC Agile S NI"),
    "BT-ET6XEX2-AND": ("bartec-et60-65ex2-android", "BARTEC ET60/65EX2 Android"),
    "BT-ET6XEX2-WIN": ("bartec-et60-65ex2-windows", "BARTEC ET60/65EX2 Windows"),
    "BT-AGILE-X-IS": ("bartec-agile-x-is", "BARTEC Agile X IS"),
    # Aegex, Eaton, TouchStar
    "AEG-100M": ("aegex-100m", "Aegex 100M"),
    "AEG-10-IS": ("aegex-10-is", "Aegex 10 IS"),
    "EA-GECMA-IS": ("eaton-gecma-is-tablet-pc", "Eaton MTL GECMA IS Tablet PC"),
    "TS-3300": ("touchstar-ts3300", "TouchStar TS3300"),
    # Getac: the current generation keeps the model URL, earlier generations get a suffix
    "GT-F110EX": ("getac-f110-ex", "Getac F110-EX G7"),
    "GT-F110EX-G6": ("getac-f110-ex-g6", "Getac F110-EX G6"),
    "GT-UX10EX": ("getac-ux10-ex", "Getac UX10-EX G5"),
    "GT-UX10EX-G2R": ("getac-ux10-ex-g2-r", "Getac UX10-EX G2-R"),
    "GT-K120EX": ("getac-k120-ex", "Getac K120-EX G3"),
    "GT-K120EX-G2R": ("getac-k120-ex-g2-r", "Getac K120-EX G2-R"),
    "GT-ZX10EX": ("getac-zx10-ex", "Getac ZX10-EX G2"),
    "GT-ZX10EX-G1": ("getac-zx10-ex-g1", "Getac ZX10-EX (first generation)"),
    "GT-ZX80EX": ("getac-zx80-ex", "Getac ZX80-EX"),
    "GT-ZX80WEX": ("getac-zx80w-ex", "Getac ZX80W-EX"),
    "GT-ZX70EX": ("getac-zx70-ex", "Getac ZX70-EX"),
    "GT-EX80": ("getac-ex80", "Getac EX80"),
    "GT-T800EX": ("getac-t800-ex", "Getac T800-Ex"),
    # Durabook, Winmate, Zebra, Panasonic, Dell
    "DB-R8EX": ("durabook-r8-ex", "Durabook R8-EX"),
    "DB-U11IEX": ("durabook-u11i-ex", "Durabook U11I-EX"),
    "WM-M101ADEX": ("winmate-m101ad-ex", "Winmate M101AD-EX"),
    "WM-M101SEX": ("winmate-m101s-ex", "Winmate M101S-EX"),
    "WM-M700DQ8EX": ("winmate-m700dq8-ex", "Winmate M700DQ8-EX"),
    "ZB-L10ATEX-WIN": ("zebra-l10-atex-windows", "Zebra L10 ATEX (Windows)"),
    "ZB-L10ATEX-AND": ("zebra-l10-atex-android", "Zebra L10 ATEX (Android)"),
    "PA-FZG1ATEX-MK4": ("panasonic-toughpad-fz-g1-atex", "Panasonic Toughpad FZ-G1 ATEX"),
    "DL-7220EX": ("dell-latitude-7220ex", "Dell Latitude 7220EX"),
    # iPad cases and conversions
    "XS-IPP11": ("xshielder-ipad-pro-11", "Xshielder iPad Pro 11-inch"),
    "XC-XCZ1-M7": ("xciel-xcz1-ipad-mini-7", "Xciel XCZ1 iPad mini 7"),
    "XC-XCZ1-M6": ("xciel-xcz1-ipad-mini-6", "Xciel XCZ1 iPad mini 6"),
    "XC-XCZ1-M5": ("xciel-xcz1-ipad-mini-5", "Xciel XCZ1 iPad mini 5"),
    "XC-XCZ1-M4": ("xciel-xcz1-ipad-mini-4", "Xciel XCZ1 iPad mini 4"),
    "XC-XCR-A16": ("xciel-xcripad-11", "Xciel XCRiPad 11-inch"),
    "XC-XCR-M7": ("xciel-xcripad-mini-7", "Xciel XCRiPad mini 7"),
    "XC-XCR-AIR11": ("xciel-xcripad-air-11", "Xciel XCRiPad Air 11-inch"),
    "XC-XCR-PRO11": ("xciel-xcripad-pro-11", "Xciel XCRiPad Pro 11-inch"),
    "JF-HTA3021": ("joy-factory-axtion-extreme-ipad-mini", "The Joy Factory aXtion Extreme iPad mini"),
    "AX-IPM7-Z1": ("atexxo-zone-1-ipad-mini-7", "Atexxo Zone 1 iPad mini 7"),
    "AX-IPM6-Z1": ("atexxo-zone-1-ipad-mini-6", "Atexxo Zone 1 iPad mini 6"),
    "AX-IPM5-Z1": ("atexxo-zone-1-ipad-mini-5", "Atexxo Zone 1 iPad mini 5"),
    "AX-IPA16-Z2": ("atexxo-zone-2-ipad-a16", "Atexxo Zone 2 iPad (A16)"),
    "AX-IPP11-Z2": ("atexxo-zone-2-ipad-pro-11", "Atexxo Zone 2 iPad Pro 11-inch"),
    "AX-IPM7-Z2": ("atexxo-zone-2-ipad-mini-7", "Atexxo Zone 2 iPad mini 7"),
    "AX-IPM6-Z2": ("atexxo-zone-2-ipad-mini-6", "Atexxo Zone 2 iPad mini 6"),
    "AX-IPAD109-Z2": ("atexxo-zone-2-ipad-10-9", "Atexxo Zone 2 iPad 10.9-inch"),
    "AR-UNICASE": ("armadex-ex-unicase", "Armadex Ex UniCase"),
    "OB-HARDLINE-IPADAIR5": ("otterbox-hardline-ipad-air", "OtterBox Hardline iPad Air"),
    "RN-AP10.2Z2.22": ("rino-ex-nr-case-ipad-10-2", "RINO Solutions Ex nR case iPad 10.2-inch"),
    # Samsung conversions
    "AX-GTA5-Z1": ("atexxo-zone-1-galaxy-tab-active5", "Atexxo Zone 1 Galaxy Tab Active5"),
    "AX-GTA5P-Z2": ("atexxo-zone-2-galaxy-tab-active5-pro", "Atexxo Zone 2 Galaxy Tab Active5 Pro"),
    "AX-S10FE-Z2": ("atexxo-zone-2-galaxy-tab-s10-fe", "Atexxo Zone 2 Galaxy Tab S10 FE"),
    "AX-S9FE-Z2": ("atexxo-zone-2-galaxy-tab-s9-fe", "Atexxo Zone 2 Galaxy Tab S9 FE"),
    "AX-A9-Z2": ("atexxo-zone-2-galaxy-tab-a9", "Atexxo Zone 2 Galaxy Tab A9"),
}

# Workbook "Brand" -> (brand, brandSlug, maker, country, website). website is None where the maker has no live site.
BRANDS = {
    "i.safe MOBILE": ("i.safe MOBILE", "isafe", "i.safe MOBILE GmbH", "Germany", "https://www.isafe-mobile.com"),
    "ecom / Pepperl+Fuchs": ("ecom (Pepperl+Fuchs)", "ecom", "ecom instruments (Pepperl+Fuchs)", "Germany", "https://www.ecom-ex.com"),
    "BARTEC": ("BARTEC", "bartec", "BARTEC GmbH", "Germany", "https://bartec.com"),
    "Aegex": ("Aegex", "aegex", "Aegex Technologies LLC", "USA", None),  # aegex.com stopped resolving (2026-10)
    "Eaton MTL GECMA": ("Eaton", "eaton", "Eaton Electric Limited (MTL)", "United Kingdom", "https://www.eaton.com"),
    "TouchStar": ("TouchStar", "touchstar", "TouchStar Technologies Ltd", "United Kingdom", "https://touchstar.co.uk"),
    "Getac": ("Getac", "getac", "Getac Technology Corporation", "Taiwan", "https://www.getac.com"),
    "Durabook": ("Durabook", "durabook", "Twinhead International Corp. (Durabook)", "Taiwan", "https://www.durabook.com"),
    "Winmate": ("Winmate", "winmate", "Winmate Inc.", "Taiwan", "https://www.winmate.com"),
    "Zebra (XSLATE / XPAD L10 ATEX)": ("Zebra", "zebra", "Zebra Technologies", "USA", "https://www.zebra.com"),
    "Panasonic Toughpad / Toughbook": ("Panasonic", "panasonic", "Panasonic Connect", "Japan", "https://eu.connect.panasonic.com"),
    "Dell Latitude": ("Dell", "dell", "Dell Technologies", "USA", "https://www.dell.com"),
    "Xshielder": ("Xshielder", "xshielder", "Xshielder AS", "Norway", "https://xshielder.com"),
    "Xciel": ("Xciel", "xciel", "Xciel Inc", "USA", "https://www.xciel.com"),
    "The Joy Factory": ("The Joy Factory", "joy-factory", "The Joy Factory Inc.", "USA", "https://thejoyfactory.com"),
    "Atexxo": ("Atexxo", "atexxo", "Atexxo B.V.", "Netherlands", "https://atexxo.com"),
    "Holland Shielding": ("Holland Shielding", "holland-shielding", "Holland Shielding Systems BV", "Netherlands", "https://hollandshielding.com"),
    "Armadex": ("Armadex", "armadex", "Armadex Explosion Protection B.V.", "Netherlands", "https://armadex.com"),
    "OtterBox": ("OtterBox", "otterbox", "Otter Products, LLC", "USA", "https://www.otterbox.com"),
    "RINO Solutions": ("RINO Solutions", "rino", "RINO Solutions Ltd", "United Kingdom", None),
}
SPONSOR_BRANDS = {"xshielder"}

# Short text for each folded version (the parent page lists it under "Versions")
VARIANT_TEXT = {
    "IS-MOP1A.1": ("IS-MOP1A.1 Mobile Operator Panel",
                   "The IS940.1 combined with the IS-CS1A.1 connectivity system as a docked operator panel. The connectivity "
                   "module needs a separate Ex enclosure."),
    "IS-MOP1B.1": ("IS-MOP1B.1 Mobile Operator Panel",
                   "The IS945.1 combined with the IS-CS1A.1 connectivity system as a docked operator panel. The connectivity "
                   "module needs a separate Ex enclosure."),
    "IS-930.M1": ("IS930.M1", "Mining version, also certified for Group I (underground mining)."),
    "IS-940.M1": ("IS940.M1", "Mining version, also certified for Group I (underground mining)."),
    "IS-945.M1": ("IS945.M1", "Mining version, also certified for Group I (underground mining)."),
    "IS-910.M1": ("IS910.M1", "Mining version, also certified for Group I (underground mining). Discontinued."),
    "PF-TE05-D2": ("Tab-Ex 05 D2", "North American version, certified for Class I Division 2 only, with no ATEX or IECEx rating."),
    "HS-TABEXD-IPADM5": ("Holland Shielding TAB-EXd-IPAD-M5",
                         "Holland Shielding sold this conversion under its own order code. The certificate names Atexxo B.V."),
}
MINING_VARIANTS = {"IS-930.M1", "IS-940.M1", "IS-945.M1", "IS-910.M1"}

# Status notes: required for every discontinued product and every flagged status cell. (status, note, successor ID)
STATUS = {
    "IS-910.1": ("discontinued", "i.safe MOBILE has moved it to the archive of its range.", "IS-930.1"),
    "IS-910.2": ("discontinued", "i.safe MOBILE has moved it to the archive of its range.", "IS-930.2"),
    "PF-TE02-DZ1": ("discontinued", "ecom lists it as discontinued.", "PF-TE03-DZ1"),
    "PF-TE02-DZ2": ("discontinued", "ecom lists it as discontinued.", "PF-TE03-DZ2"),
    "PF-TE03-DZ1": ("discontinued", "ecom lists it as discontinued. Security patches ended in Q3 2025; hardware service runs to Q4 2027.", "PF-TE05-DZ1"),
    "PF-TE03-DZ2": ("discontinued", "ecom lists it as discontinued. Security patches ended in Q3 2025; hardware service runs to Q4 2027.", "PF-TE05-DZ2"),
    "PF-TEPRO-DZ2": ("discontinued", "ecom lists it as discontinued.", "PF-TE04PRO-DZ2"),
    "PF-TE04PRO-DZ2": ("discontinued", "ecom lists it as discontinued and names no 10-inch successor.", None),
    "PF-PADEX01-P8": ("discontinued", "ecom lists it as discontinued.", "PF-PADEX01-P12-DZ2"),
    "BT-AGILE-X-IS": ("discontinued", "BARTEC ended sales in April 2023 and support in April 2026. It names no Zone 1 successor.", None),
    "AEG-100M": ("discontinued", "Availability unclear: Aegex's website has gone offline and a distributor lists it as end of life, "
                                 "though another still shows stock. Eaton sells a tablet with identical specifications.", "EA-GECMA-IS"),
    "AEG-10-IS": ("discontinued", "Discontinued by Aegex; its IECEx certificate has been cancelled.", "AEG-100M"),
    "GT-F110EX-G6": ("discontinued", "Getac now lists the G7 generation in its place; it published no end-of-sale date.", "GT-F110EX"),
    "GT-UX10EX-G2R": ("discontinued", "Getac now lists the G5 generation in its place; it published no end-of-sale date.", "GT-UX10EX"),
    "GT-K120EX-G2R": ("discontinued", "Getac now lists the G3 generation in its place; it published no end-of-sale date.", "GT-K120EX"),
    "GT-ZX10EX-G1": ("discontinued", "Getac now lists the G2 generation in its place; it published no end-of-sale date.", "GT-ZX10EX"),
    "GT-ZX70EX": ("discontinued", "No longer in Getac's range, which now lists the ZX80-EX; no end-of-sale notice.", "GT-ZX80EX"),
    "GT-EX80": ("discontinued", "No longer in Getac's range; Getac names no Zone 0/1 successor.", None),
    "GT-T800EX": ("discontinued", "No longer in Getac's range; no end-of-sale notice.", None),
    "WM-M700DQ8EX": ("discontinued", "Winmate lists it as discontinued. Its successor, the M700DQ9, has no Ex version.", None),
    "ZB-L10ATEX-WIN": ("discontinued", "Zebra no longer sells the L10 Windows range; its ATEX versions aren't in Zebra's current range.", None),
    "ZB-L10ATEX-AND": ("discontinued", "No longer in Zebra's range. The IECEx certificate for this version was suspended on 6 March 2026; "
                                       "units made before then are unaffected.", None),
    "PA-FZG1ATEX-MK4": ("discontinued", "No longer sold. Panasonic's later Toughbook G2 and G3 have no ATEX version.", None),
    "DL-7220EX": ("discontinued", "No longer sold, and its IECEx certificate was cancelled in 2023. Dell has no Ex successor.", None),
    "XC-XCZ1-M5": ("current", "Xciel still lists it and says it has iPad mini 5 stock. Apple no longer sells the iPad mini 5, so supply is limited.", None),
    "XC-XCZ1-M4": ("discontinued", "No longer in Xciel's range.", "XC-XCZ1-M5"),
    "AX-IPM5-Z1": ("discontinued", "No longer in Atexxo's range.", "AX-IPM6-Z1"),
    "AR-UNICASE": ("current", "New in 2026. Armadex lists it for pre-order and also says it ships from stock; check lead times when ordering.", None),
    "OB-HARDLINE-IPADAIR5": ("discontinued", "Its IECEx certificate was cancelled in March 2025 and OtterBox no longer promotes it; "
                                              "we found no end-of-sale notice.", None),
    "RN-AP10.2Z2.22": ("discontinued", "Its IECEx certificate was cancelled in 2025.", None),
    **{pid: ("current", "Listed by Atexxo; one distributor shows it as out of stock.", None)
       for pid in ("AX-GTA5P-Z2", "AX-S10FE-Z2", "AX-S9FE-Z2", "AX-A9-Z2")},
}

# ---------------------------------------------------------------------------------------------------------------
# Public notes for flagged (CHECK / CONFLICT) cells that are published.
# Rules are tried in order: (sheet, record glob, fields, note). fields is a set of column names, or None for any.
# Drafts for Rory's review (2026-10-26): see ../public-note-templates.md
# ---------------------------------------------------------------------------------------------------------------

CERT_FIELDS = {"Marking - gas", "Marking - dust", "Marking - mining", "Protection concepts", "EPL gas", "EPL dust", "Gas group",
               "Dust group", "Temp class", "Max surface temp dust", "Ex ambient Ta", "Certificate status", "ATEX certificate",
               "IECEx certificate", "Specific conditions / safety rules", "Ex IP (cert)"}

N_CLAIM_NO_CERT = ("Marking as stated by {brand}. {brand} doesn't publish the certificate or declaration of conformity, "
                   "so we haven't been able to check it.")
N_GETAC_MAP = ("Getac's certificate names the model with a 'c' suffix and doesn't name the generation. We matched it to "
               "this generation by its batteries, dust group and ambient range.")
N_GETAC_DUST = ("Units built under Getac's earlier certificate are rated for dust group IIIB; later units are IIIC. "
                "Check the label on the device.")
N_FAMILY_CERT = "From the certificate for the {family} range, which doesn't name this exact model."
N_TA_DEFAULT = "The certificate marks no ambient range, so the standard default applies (−20 °C to +40 °C, IEC 60079-0)."
N_ATEX_FROM_SOURCE = "The IECEx marking is from the certificate. The ATEX version is from {source}; the ATEX certificate isn't published."
N_NA = "From {source}; we haven't seen the North American certificate."
N_ASSUMED_SILENT = "Not stated in the certificate. {reason}"
N_MAKER_ONLY = "Stated by {brand}; not in the certificate conditions."

NOTE_RULES = [
    # BARTEC's own sources disagree (product page IP64, datasheet IP66); IP64 is the figure in the Ex marking
    ("Ex products", "BT-ET6XEX2-AND", {"IP rating"},
     "BARTEC's product page gives IP64 and its datasheet IP66; IP64 is the figure in the Ex marking."),
    # The Xshielder iPad cover's certification cells carry no flag: Rory confirmed it's ATEX certified and asked for it
    # to be stated plainly (decision D3, 2026-10-26; logged on the workbook's Fact check sheet).
    # Maker claims with no certificate or declaration published
    ("Ex products", "AX-*-Z2", CERT_FIELDS, N_CLAIM_NO_CERT),
    ("Ex products", "XC-XCR-*", CERT_FIELDS, N_CLAIM_NO_CERT),
    ("Ex products", "AX-*-Z2", {"Ex-area restrictions on peripherals", "Other approvals"}, N_MAKER_ONLY),
    ("Ex products", "AX-GTA5-Z1", {"Ex-area restrictions on peripherals"},
     "The SIM and cover rules are from the certificate. Atexxo's page adds a microSD rule that isn't in the certificate."),
    ("Ex products", "AX-GTA5-Z1", {"Other approvals"}, "Samsung's MIL-STD rating for the tablet itself; Atexxo doesn't say whether it holds in the Ex housing."),
    # Getac certificate mapping and dust groups
    ("Ex products", "GT-F110EX-G6", {"Marking - dust", "Dust group"}, N_GETAC_DUST),
    ("Ex products", "GT-UX10EX-G2R", {"Marking - dust", "Dust group", "Ex ambient Ta"}, N_GETAC_DUST),
    ("Ex products", "GT-ZX80WEX", {"Ex ambient Ta", "Specific conditions / safety rules", "IECEx certificate"},
     "No certificate names the Windows ZX80W-EX yet. This is from the ZX80-EX certificate, which covers the same hardware."),
    ("Ex products", "GT-*", {"IECEx certificate"}, N_GETAC_MAP),
    ("Ex products", "GT-ZX70EX", {"Charging in Ex area", "Battery swap in Ex area"},
     N_ASSUMED_SILENT.format(reason="The battery isn't user-removable, and later Getac Ex tablets forbid charging in the zone.")),
    ("Ex products", "GT-ZX70EX", {"Ex-area restrictions on peripherals"}, "Not stated in the certificate; based on the rubber I/O cover it lists as an accessory."),
    ("Ex products", "GT-EX80", {"Marking - gas", "Marking - dust"}, N_ATEX_FROM_SOURCE.format(source="Getac's own information")),
    # Single-family certificates
    ("Ex products", "WM-M700DQ8EX", {"Marking - gas", "Temp class", "Ex ambient Ta", "IECEx certificate"},
     N_FAMILY_CERT.format(family="Winmate M700")),
    ("Ex products", "ZB-L10ATEX-WIN", {"IECEx certificate", "Certificate status"},
     "The IECEx certificate for the Windows version isn't in the IECEx database; the related Android certificate was suspended on 6 March 2026."),
    ("Ex products", "ZB-L10ATEX-*", {"Specific conditions / safety rules", "Ex ambient Ta"},
     "From the IECEx certificate for the Android version (suspended in March 2026); the Windows version's certificate isn't published."),
    # ATEX parts not read at source
    ("Ex products", "AEG-100M", {"Marking - gas", "Marking - dust"}, N_ATEX_FROM_SOURCE.format(source="a distributor and Eaton's identical certificate")),
    ("Ex products", "XC-XCZ1-*", {"Marking - gas", "ATEX certificate"}, N_ATEX_FROM_SOURCE.format(source="a distributor")),
    ("Ex products", "JF-HTA3021", {"Marking - gas", "Marking - dust"}, N_ATEX_FROM_SOURCE.format(source="The Joy Factory's datasheet")),
    ("Ex products", "TS-3300", {"Marking - gas", "Marking - dust"},
     "Only an IECEx certificate is published. The ATEX marking is from TouchStar's datasheet."),
    ("Ex products", "DL-7220EX", {"Marking - gas", "Marking - dust"}, "From the IECEx certificate (cancelled in 2023); Dell's ATEX certificate isn't published."),
    ("Ex products", "DL-7220EX", {"Ex ambient Ta"}, "Stated in the first issue of the IECEx certificate only; later issues don't repeat it."),
    ("Ex products", "DB-*", {"ATEX certificate"}, "From Durabook's user manual, which reproduces the certificate of conformity; the certificate itself isn't published."),
    ("Ex products", "EA-GECMA-IS", {"Ex IP (cert)"}, "From Eaton's manual; the IECEx certificate doesn't state an IP code."),
    ("Ex products", "BT-AGILE-X-IS", {"Marking - mining"}, "From BARTEC's quick start guide; the IECEx database also lists an Mb type."),
    ("Ex products", "AEG-10-IS", {"Temp class"}, "The certificate allows T6, T4 or T3 depending on the build; check the label on the device."),
    ("Ex products", "OB-HARDLINE-*", {"Temp class"}, "OtterBox's quick start guide says T5; the IECEx certificate (now cancelled) said T6."),
    # Default ambient range when the certificate marks none
    ("Ex products", "JF-HTA3021", {"Ex ambient Ta"}, N_TA_DEFAULT),
    ("Ex products", "RN-*", {"Ex ambient Ta"}, N_TA_DEFAULT),
    ("Ex products", "IS-945.1", {"Ex-area restrictions on peripherals"}, "From the certificate shared with the IS940.1."),
    # Battery swap: certain for iPads (battery not user-replaceable), assumed elsewhere
    ("Ex products", "XC-*", {"Battery swap in Ex area"}, "The iPad's battery isn't user-replaceable."),
    ("Ex products", "AX-IP*", {"Battery swap in Ex area"}, "The iPad's battery isn't user-replaceable."),
    ("Ex products", "PA-FZG1ATEX-MK4", {"Battery swap in Ex area"},
     N_ASSUMED_SILENT.format(reason="It forbids all external connections in the zone.")),
    ("Ex products", "DL-7220EX", {"Charging in Ex area"}, "The certificate forbids connecting external devices in the zone, which rules out charging."),
    # Status cells (the public text is the status note)
    ("Ex products", "*", {"Status"}, "{status_note}"),
    # Base devices
    ("Base devices", "B-APL-*", {"RAM (GB)"}, "Apple doesn't publish RAM; from independent teardowns."),
    ("Base devices", "B-APL-IPADPRO*", {"Security updates until"},
     "Apple doesn't publish an end date; its legal minimum is five years from first sale. Previous iPad Pro models have had "
     "security updates for seven to eight years."),
    ("Base devices", "B-APL-*", {"Security updates until"}, "Apple's legal minimum (five years from first sale). Earlier iPads have had longer support."),
    ("Base devices", "B-SAM-TABA9", {"Security updates until"}, "Samsung publishes no date for this model; this is the Galaxy Tab A9+ date."),
    ("Base devices", "B-SAM-TABACTIVE5", {"Security updates until"},
     "Samsung's date for the standard model. The Enterprise Edition gets longer support; the converter doesn't say which it uses."),
    ("Base devices", "B-GT-*", {"Security updates until"}, "Getac's date for the standard model; Getac lists no separate date for the EX version."),
    ("Base devices", "B-EC-TE05", {"Security updates until"}, "ecom states security updates until 2032, which matches Samsung's Enterprise Edition commitment."),
    *[("Base devices", b, {"Security updates until"},
       "Microsoft ended Windows 10 Pro support on 14 October 2025. The maker doesn't state the edition; an LTSC edition would run longer.")
      for b in ("B-AEG-10", "B-PN-FZG1-MK4", "B-EC-PADEX-P8", "B-ZB-L10-WIN")],
    *[("Base devices", b, {"Security updates until"},
       "Microsoft's end date for Windows 11 IoT Enterprise LTSC 2024. The maker doesn't state the edition year.")
      for b in ("B-IS-945", "B-AEG-100M", "B-WM-M101AD", "B-GT-ZX80W")],
    ("Base devices", "B-BT-AGILE-X-IS", {"Security updates until"}, "BARTEC's own end of support for this tablet."),
    ("Base devices", "B-BT-ET60EX2-AND", {"Security updates until"},
     "Zebra's security-update date for the ET60/ET65 the tablet is built on; BARTEC says it needs a LifeGuard service contract."),
    ("Base devices", "*", {"Battery (Wh)"}, "Calculated from the maker's battery capacity and voltage; the maker doesn't state watt-hours."),
    ("Base devices", "B-APL-*", {"Battery (mAh)"}, "Apple states battery capacity in watt-hours only; this figure is from the case maker's documents."),
    ("Base devices", "*", {"OS edition"}, "The maker doesn't state the edition year."),
]

# Generic notes by field, used when no rule above matches. Certification fields must match a rule above, or the
# import stops; these are for lower-stakes fields.
FIELD_NOTES = {
    ("Ex products", "Launch"): "Approximate: the maker doesn't publish a launch date. Based on the certificate date or the maker's announcement.",
    ("Ex products", "End of sale / EOL"): "Approximate: the maker published no end-of-sale date.",
    ("Ex products", "Successor"): "The maker hasn't named a successor; this is the model that took its place in the maker's range.",
    ("Ex products", "Certified accessories"): "Not checked against the certificate for this exact model.",
    ("Ex products", "Tablet removable and reusable"): "Not stated by the maker for this exact model.",
    ("Ex products", "Fitting"): "Not stated by the maker for this exact model.",
    ("Ex products", "North America (NEC/CEC)"): N_NA.format(source="the maker's published information"),
    ("Ex products", "Other approvals"): "Stated by the maker; we haven't seen the certificate.",
    ("Ex products", "UKEX certificate"): "We haven't found the UK certificate itself.",
    ("Ex products", "In-use weight (g)"): "From a single source; the maker's figures vary.",
    ("Ex products", "Dimensions (mm)"): "From a single source.",
    ("Ex products", "IP rating"): "From a single source; not stated in the Ex certificate.",
    ("Ex products", "Operating temp (°C)"): "From a single source.",
    ("Ex products", "Drop / MIL-STD"): "From a single source.",
    ("Ex products", "Camera-free variant"): "Listed by a distributor; the maker doesn't list it.",
    ("Ex products", "Part numbers / SKUs"): "Order codes from a certificate or distributor; confirm them with the maker when ordering.",
    ("Ex products", "Charging in Ex area"): N_ASSUMED_SILENT.format(reason="Assumed from the maker's instructions."),
    ("Ex products", "Battery swap in Ex area"): N_ASSUMED_SILENT.format(reason="Assumed from the maker's instructions."),
    ("Ex products", "Ex-area restrictions on peripherals"): "Not stated in the certificate for this exact model.",
    ("Base devices", "First supply date"): "Approximate.",
    ("Base devices", "*"): "From a single or secondary source; the maker doesn't state it for this model.",
    ("Through-case access", "*"): "Inferred from related information; not stated by the maker for this exact model.",
    ("Capabilities", "*"): "Not confirmed by the maker for this exact model.",
}

# Flagged cells that are never published. (sheet, record glob, field, reason)
SKIP_RULES = [
    ("Ex products", "AX-*-Z2", "North America (NEC/CEC)", "marketing equivalence, not a North American approval"),
    ("Ex products", "XC-XCR-*", "North America (NEC/CEC)", "Xciel claim with no certificate; North America is out of scope"),
    ("Ex products", "XC-XCZ1-*", "North America (NEC/CEC)", "the datasheet lists gas groups, not a North American approval"),
    ("Ex products", "AX-GTA5-Z1", "North America (NEC/CEC)", "'on request' only"),
    ("Ex products", "AX-IPM*-Z1", "UKEX certificate", "UKCA mentioned on a product page; no UK certificate"),
    ("Ex products", "AX-*", "Marking - mining", "no Group I marking shown"),
    ("Ex products", "XS-IPP11", "Specific conditions / safety rules", "the certificate's conditions aren't published yet"),
]

SECURITY_QUARTERS = {"B-IS-930", "B-IS-940", "B-IS-945", "B-IS-941", "B-IS-910", "B-EC-TE05", "B-EC-TE03", "B-EC-TE04PRO",
                     "B-EC-PADEX-P12"}
SOURCE_ORDER = ["Certificate", "DoC", "Database", "Datasheet", "Manual", "Maker page", "Maker statement", "Distributor", "Press"]

THROUGH_LABELS = {
    "Biometric unlock (Face ID / Touch ID)": "Face ID / Touch ID",
    "Apple Pencil / stylus": "Apple Pencil / stylus",
    "Smart Connector / keyboard": "Smart Connector / keyboard",
    "Side and volume buttons": "Side and volume buttons",
    "Rear cameras unobstructed": "Rear cameras",
    "LiDAR unobstructed": "LiDAR",
    "Front camera": "Front camera",
    "Touchscreen": "Touchscreen",
    "Screen cover": "Screen cover",
    "NFC through case": "NFC",
    "Charging method": "Charging",
    "Speakers and microphones": "Speakers and microphones",
    "Headset audio": "Headset audio",
    "USB-C access": "USB-C port",
    "SIM access": "SIM card",
    "Removed or disabled components": "Removed or disabled",
    "Mounting points": "Mounting points",
}
CAPABILITY_LABELS = {
    "PTT": "Push-to-talk",
    "Emergency / SOS key": "Emergency / SOS",
    "Lone worker": "Lone worker",
    "Barcode scanning": "Barcode scanning",
    "RFID": "RFID",
    "Thermal imaging": "Thermal imaging",
    "NFC reading": "NFC reading",
    "3D scanning / LiDAR": "3D scanning / LiDAR",
    "High-accuracy GNSS": "High-accuracy GNSS",
    "Satellite": "Satellite",
    "Remote expert / video calls": "Remote expert / video calls",
    "Windows desktop software": "Windows desktop software",
    "Bluetooth gas detector integration": "Bluetooth gas detectors",
}

# ---------------------------------------------------------------------------------------------------------------
# Workbook access
# ---------------------------------------------------------------------------------------------------------------

MISSING = []  # (sheet, id, column, value, comment) for flagged cells that would be published without a public note
NOTES_USED = []  # (sheet, id, column, note) for the report
SKIPPED = []  # (sheet, id, column, reason)


class Row:
    def __init__(self, sheet, cells, key):
        self.sheet = sheet
        self.cells = cells
        self.id = cells[key].v

    def raw(self, col):
        c = self.cells.get(col)
        return c.v if c else None

    def flag(self, col):
        c = self.cells.get(col)
        return c.flag if c else None

    def flagged(self, col):
        return self.flag(col) in ("check", "conflict")

    def skip_reason(self, col):
        for sheet, pattern, field, reason in SKIP_RULES:
            if sheet == self.sheet and field == col and fnmatch(self.id, pattern) and self.flagged(col):
                return reason
        return None

    def get(self, col):
        """Publishable value: None when empty, NOT FOUND, a placeholder, a formula or skipped."""
        c = self.cells.get(col)
        if c is None or c.v is None or c.flag == "notfound":
            return None
        if isinstance(c.v, str) and (PLACEHOLDER.match(c.v) or c.v.startswith("=")):
            return None
        reason = self.skip_reason(col)
        if reason:
            SKIPPED.append((self.sheet, self.id, col, reason))
            return None
        return c.v


def note_for(row, col, ctx):
    for sheet, pattern, fields, text in NOTE_RULES:
        if sheet == row.sheet and fnmatch(row.id, pattern) and (fields is None or col in fields):
            return text.format(**ctx)
    if row.sheet == "Ex products" and col in CERT_FIELDS:
        return None  # certification fields need an explicit rule
    return FIELD_NOTES.get((row.sheet, col)) or FIELD_NOTES.get((row.sheet, "*"))


def use(row, col, key, notes, ctx=None):
    """Value of a cell, recording a public note when the cell is flagged."""
    v = row.get(col)
    if v is not None and row.flagged(col):
        text = note_for(row, col, ctx or {})
        if not text:
            MISSING.append((row.sheet, row.id, col, v, row.cells[col].comment))
        else:
            NOTES_USED.append((row.sheet, row.id, col, text))
            if key:
                notes[key] = text
    return v


def load(path):
    wb = openpyxl.load_workbook(path)

    def sheet(name, key):
        ws = wb[name]
        header = [c.value for c in ws[1]]
        rows = []
        for r in ws.iter_rows(min_row=2):
            if r[0].value is None:
                continue
            cells = {}
            for c in r:
                if c.column > len(header) or header[c.column - 1] is None:
                    continue
                v = c.value
                if isinstance(v, datetime.datetime):
                    v = v.date().isoformat()
                elif isinstance(v, str):
                    v = v.strip()
                fill = (c.fill.fgColor.rgb or "")[-6:] if c.fill is not None and c.fill.fill_type == "solid" else ""
                cells[header[c.column - 1]] = Cell(v, FLAG_FILLS.get(str(fill).upper()), c.comment.text if c.comment else "")
            rows.append(Row(name, cells, key))
        return {row.id: row for row in rows}

    return (
        sheet("Ex products", "Product ID"),
        sheet("Base devices", "Base device ID"),
        sheet("Through-case access", "Product ID"),
        sheet("Capabilities", "Product ID"),
        sheet("Sources", "ID"),
    )


# ---------------------------------------------------------------------------------------------------------------
# Value helpers
# ---------------------------------------------------------------------------------------------------------------

def minus(text):
    """Typographic minus for negative temperatures only ('-20 °C', '-20 to'), not for standard numbers"""
    return re.sub(r"(?<![\w./])-(?=\d+(?:\s*°C|\s*C\b|\s+to\b))", "−", text)


def temp_range(text):
    """'-20 to +55' or '-20 °C to +55 °C' -> ('−20 °C to +55 °C', -20, 55)"""
    if text is None:
        return None, None, None
    text = re.sub(r"\s*operating\b", "", str(text).split(";")[0])
    m = re.match(r"^\s*([+-]?\d+)\s*(?:°\s*C|C)?\s*(?:to|…|\.\.\.|–)\s*([+-]?\d+)\s*(?:°\s*C|C)?\s*(.*)$", text.replace("−", "-"))
    if not m:
        return minus(text), None, None
    lo, hi, rest = int(m.group(1)), int(m.group(2)), m.group(3)
    fmt = lambda n: f"{'+' if n > 0 else ''}{n}".replace("-", "−")
    return f"{fmt(lo)} °C to {fmt(hi)} °C" + (f" {minus(rest)}" if rest else ""), lo, hi


def number(v):
    if isinstance(v, (int, float)):
        return int(v)
    if isinstance(v, str):
        m = re.match(r"^\s*(\d[\d,]*)\s*(?:g\b|$)", v)
        if m:
            return int(m.group(1).replace(",", ""))
    return None


def first_number(v):
    m = re.search(r"\d+(?:\.\d+)?", str(v or ""))
    return float(m.group(0)) if m else None


def split_scope(text):
    m = re.match(r"^(.*?)\s*\((.*)\)\s*$", text)
    return (m.group(1).strip(), m.group(2).strip()) if m else (text.strip(), None)


MARK_RE = re.compile(
    r"(?:(?P<grp>I{1,2})\s?(?P<cat>[123])\s?(?P<med>GD|G|D)\b\s*;?\s*)?"
    r"Ex\s+(?P<conc>(?:(?:i[abc]|d[abc]|e[bc]|t[abc]|m[abc]|n[ARC]|op|is|pr|sh)\s+)+)"
    r"(?P<group>III[ABC]|II[ABC]|I)?\s*(?P<temp>T[1-6](?![\d°])|T?\d+\s?°C)?\s*(?P<epl>[GDM][abc])\b(?:\s+(?P<ip>IP\d[\dXx]))?")


def find_markings(text):
    """Every Ex marking in a cell: [{'atex': has an ATEX category, 'full': with it, 'plain': without, 'temp': …}]"""
    if not text:
        return []
    text = re.sub(r"\((?:[^()]*\b(?:assumed|not read)\b[^()]*)\)", "", str(text))  # asides about unread markings
    out = []
    for m in MARK_RE.finditer(text):
        body = " ".join(x for x in ("Ex", m["conc"].strip(), m["group"], (m["temp"] or "").replace(" ", ""), m["epl"],
                                    (m["ip"] or "").upper()) if x)
        prefix = f"{m['grp']} {m['cat']}{m['med']}" if m["cat"] else None
        out.append({"atex": bool(prefix), "full": f"{prefix} {body}" if prefix else body, "plain": body, "temp": m["temp"]})
    return out


def clean_marking_text(text):
    """A marking cell for the Versions list, without the research shorthand"""
    if not text or not find_markings(text):
        return None
    t = re.sub(r"\s*\((?:ATEX|IECEx|G\d|R\d)[^()]*\)", "", str(text))
    t = re.sub(r"\((?:[^()]*\b(?:assumed|not read)\b[^()]*)\)", "", t)
    return re.sub(r"\s+", " ", t).strip(" ;")


def strip_atex_prefix(marking):
    return " / ".join(re.sub(r"^I{1,2}\s+(?:[123]\s?[GD]{1,2}|M[12])\s+", "", part.strip()) for part in marking.split(" / "))


def clean(text):
    return re.sub(r"\s+-\s+", "; ", str(text)).strip() if text is not None else None


def tidy(text):
    """Through-case and capability values without research shorthand ('(datasheet …)', 'per …')"""
    t = clean(text)
    t = re.sub(r"\s*\((?:datasheet[^()]*|stated|per [^()]*|flyer[^()]*|research[^()]*)\)", "", t)
    return t.strip()


def iso_date(v):
    m = re.match(r"(\d{4}-\d{2}-\d{2})", str(v or ""))
    return m.group(1) if m else None


TODAY = datetime.date.today().isoformat()
SOON = (datetime.date.today() + datetime.timedelta(days=2 * 365)).isoformat()


def month_year(iso):
    return datetime.date.fromisoformat(iso).strftime("%B %Y")


def quarter_label(iso):
    d = datetime.date.fromisoformat(iso)
    q = (d.month - 1) // 3 + 1
    end = datetime.date(d.year + (q == 4), 1 if q == 4 else 3 * q + 1, 1) - datetime.timedelta(days=1)
    return f"Q{q} {d.year}" if d == end else None


def security_end(o):
    return o.get("securityUntilLabel") or month_year(o["securityUntil"])


IPAD_PRO_SUPPORT_YEARS = 7


def security_ended(o):
    return bool(o.get("securityUntil")) and not o.get("modernLifecycle") and not o.get("supportYears") and o["securityUntil"] < TODAY


def concepts(marking):
    tokens = re.split(r"\s+", marking)
    if "Ex" not in tokens:
        return []
    out = []
    for t in tokens[tokens.index("Ex") + 1:]:
        if not re.fullmatch(r"i[abc]|d[abc]|e[bc]|t[abc]|m[abc]|n[ARC]|op|is|pr|sh", t):
            break
        out.append(t)
    return out


def protection(marking):
    """(label, kind) from the protection concepts in the gas marking"""
    cs = concepts(marking or "")
    main = [c for c in cs if c not in ("op", "is", "pr", "sh")]
    if not main:
        return None, None
    if main[0].startswith("d"):
        extra = f" with an intrinsically safe circuit (Ex {' '.join(main[1:])})" if any(c.startswith("i") for c in main[1:]) else ""
        return f"Flameproof enclosure (Ex {main[0]}){extra}", "flameproof"
    if main[0] == "nR":
        return "Restricted-breathing enclosure (Ex nR)", "restricted"
    if main[0] in ("ec", "eb") and len(main) == 1:
        return f"Increased safety (Ex {main[0]})", "increased"
    if "ec" in main and any(c.startswith("i") for c in main):
        return f"Increased safety and intrinsic safety (Ex {' '.join(main)})", "combined"
    if main[0].startswith("i"):
        extra = " with inherently safe optical radiation (op is)" if "op" in cs else ""
        return f"Intrinsic safety (Ex {main[0]}){extra}", "intrinsic"
    return None, None


def best_epl(text, letter):
    """Highest protection level among 'Ga', 'Gb', 'Gc' found in the text"""
    found = re.findall(rf"\b{letter}([abc])\b", str(text or ""))
    return f"{letter}{min(found)}" if found else None


def zones_from_epl(epl, kind):
    table = {"gas": {"Ga": ["0", "1", "2"], "Gb": ["1", "2"], "Gc": ["2"]},
             "dust": {"Da": ["20", "21", "22"], "Db": ["21", "22"], "Dc": ["22"]}}
    return table[kind].get(epl or "", [])


def zone_label(gas, dust):
    """'Zone 1/21' for a Zone 0/1/2 + 20/21/22 product: Zone 0 isn't headlined (Rory, 2026-10-25)"""
    g, d = sorted(gas), sorted(dust)
    if len(g) > 1:
        g = [z for z in g if z not in ("0", "2")] or g[:1]
    if len(d) > 1:
        d = [z for z in d if z not in ("20", "22")] or d[:1]
    return "Zone " + "/".join(g + d[:1])


def host_name(name):
    """'iPad mini (A17 Pro), also called iPad mini 7' -> 'iPad mini (A17 Pro)'; drops Samsung model codes"""
    name = re.sub(r"\s*\((?:SM-|[^)]*SM-)[^)]*\)", "", str(name))
    depth = 0
    for i, ch in enumerate(name):
        depth += {"(": 1, ")": -1}.get(ch, 0)
        if ch == "," and depth == 0:
            return name[:i].strip()
    return name.strip()


def join_and(items):
    items = [i for i in items if i]
    return items[0] if len(items) == 1 else ", ".join(items[:-1]) + " and " + items[-1] if items else ""


def article(word):
    return "an" if re.match(r"[aeiouAEIOU]|Ex\b|8\b|11\b", word) else "a"


def source_ids(*texts):
    ids = []
    for t in texts:
        for sid in re.findall(r"\bS\d{3}\b", str(t or "")):
            if sid not in ids:
                ids.append(sid)
    return ids


# ---------------------------------------------------------------------------------------------------------------
# Building a tablet
# ---------------------------------------------------------------------------------------------------------------

TYPE_MAP = {"Ex tablet": "dedicated", "Rugged Ex variant": "rugged-variant", "Ex case": "ex-case", "Converted": "conversion"}


def build(p, bases, through, caps, sources, variants_by_parent, existing):
    pid = p.id
    slug, name = PRODUCTS[pid]
    brand_name, brand_slug, maker, country, website = BRANDS[p.raw("Brand")]
    ttype = TYPE_MAP[p.raw("Type")]
    base_ids = [b.strip() for b in str(p.raw("Base device IDs") or "").split(",") if b.strip()]
    base = bases[base_ids[0]]
    notes = {}

    # ----- Status
    status, status_note, successor = "current", None, None
    if pid in STATUS:
        status, status_note, successor = STATUS[pid]
    elif p.flagged("Status") or not str(p.raw("Status") or "").startswith("Current"):
        MISSING.append(("Ex products", pid, "Status", p.raw("Status"), "add a STATUS entry"))
    ctx = {"brand": brand_name, "status_note": status_note or ""}
    if p.flagged("Status"):
        use(p, "Status", None, notes, ctx)

    # ----- Certification
    cert = {"otherSchemes": [], "standards": []}
    gas = use(p, "Marking - gas", "certification.atex.marking", notes, ctx)
    dust = use(p, "Marking - dust", "certification.atex.marking", notes, ctx)
    gas, dust = (minus(str(gas)) if gas else None), (minus(str(dust)) if dust else None)

    # Markings are written in several ways ("ATEX: … | IECEx: …", "… (ATEX per datasheet); … (IECEx)", "with scan
    # engine: …; without: …"), so every Ex marking is found by pattern. One with an ATEX category ("II 2G") is the ATEX
    # marking; one without is the IECEx marking. Where there are several versions, the primary one matches the
    # certified temperature class and the full text is listed under Versions.
    gas_ms, dust_ms = find_markings(gas), find_markings(dust)
    tc_pref = (re.search(r"T[1-6]", str(p.raw("Temp class") or "")) or [None])[0]
    has_iecex = "IECEx" in str(p.get("IECEx certificate") or "")
    mentions_atex = "ATEX" in f"{gas} {dust} {p.raw('ATEX certificate')} {p.raw('Certificate status')}"

    def chosen(ms, atex, temp=None):
        pool = [m for m in ms if m["atex"] == atex]
        return next((m for m in pool if temp and (m["temp"] or "").startswith(temp)), pool[0] if pool else None)

    def scheme_marking(ms, temp=None):
        a, i = chosen(ms, True, temp), chosen(ms, False, temp)
        # As on the phones site: a marking with no IECEx certificate behind it is an ATEX claim (the only other scheme)
        atex_m = a["full"] if a else (i["plain"] if i and (mentions_atex or not has_iecex) else None)
        iecex_m = i["plain"] if i else (a["plain"] if a else None)
        return atex_m, iecex_m

    ga, gi = scheme_marking(gas_ms, tc_pref)
    da, di = scheme_marking(dust_ms)
    # Option lists (e.g. Aegex 10: T6, T4 or T3 builds): the dust marking follows the gas option that was picked
    plain_g = [m for m in gas_ms if not m["atex"]]
    plain_d = [m for m in dust_ms if not m["atex"]]
    if len(plain_g) > 1 and len(plain_g) == len(plain_d) and gi in [m["plain"] for m in plain_g]:
        di = plain_d[[m["plain"] for m in plain_g].index(gi)]["plain"]
        da = da if any(m["atex"] for m in dust_ms) else (di if mentions_atex else None)
    atex_marking = " / ".join(x for x in (ga, da) if x)
    iecex_marking = " / ".join(x for x in (gi, di) if x)
    if len({m["plain"] for m in gas_ms}) > 1 or len({m["plain"] for m in dust_ms}) > 1:
        shown = " | ".join(x for x in (clean_marking_text(gas), clean_marking_text(dust)) if x)
        extra_variants = [{"name": "Marking versions", "description": f"The certification covers more than one version: {shown}"}]
    else:
        extra_variants = []
    atex_raw = use(p, "ATEX certificate", "certification.atex.number", notes, ctx)
    iecex_raw = use(p, "IECEx certificate", "certification.iecex.number", notes, ctx)
    status_text = use(p, "Certificate status", "certification.status", notes, ctx)
    if atex_marking:
        atex = {"marking": atex_marking}
        a = str(atex_raw or "")
        if re.search(r"unit assessment", a, re.I):
            atex["kind"] = "unit-assessment"
        elif re.search(r"conformity|CoC", a, re.I):
            atex["kind"] = "conformity"
        elif not a or re.search(r"not (?:yet )?published|issued \(", a, re.I):
            atex["kind"] = "not-published"
        else:
            atex["kind"] = "type-examination"
        # The whole number, spaced or not: "EPS 24 ATEX 1 228 X", "Sira 19 ATEX 1017X", "TPS 23 ATEX 123637 0002 X"
        num = re.search(r"\b((?:[A-Z][A-Za-zÜ]*|TÜV|TUV)\s?\d{2}\s?ATEX\s?\d+(?:\s\d+)*(?:\s?X)?)(?!\w)", a)
        if num:
            atex["number"] = num.group(1).strip()
        if status_text:
            atex["status"] = str(status_text)
        cert["atex"] = atex
    if iecex_raw and re.search(r"IECEx", str(iecex_raw)) and iecex_marking:
        number_, scope = split_scope(re.split(r";", str(iecex_raw))[0])
        cert["iecex"] = {"marking": iecex_marking, "number": number_}
        if scope:
            cert["iecex"]["scope"] = scope[0].upper() + scope[1:]
        if "certification.atex.marking" in notes:
            notes["certification.iecex.marking"] = notes["certification.atex.marking"]
    ukex = use(p, "UKEX certificate", "certification.ukex", notes, ctx)
    if ukex and "UKEX" in str(ukex):
        cert["ukex"] = str(ukex)
    approvals = use(p, "Other approvals", "certification.otherSchemes", notes, ctx)
    if approvals:
        cert["otherSchemes"] += [a.strip() for a in re.split(r";", str(approvals)) if a.strip()]
    na = use(p, "North America (NEC/CEC)", "certification.northAmerica", notes, ctx)
    if na and not re.match(r"^(not\b|none\b|no\b|on request)", str(na), re.I):
        cert["northAmerica"] = minus(str(na))
        # The Class I (gas) Division only: "Class II, Div 2, Class III, Div 1" must not read as Class I Division 1
        div = re.search(r"Class I(?!I)\b[^;.]*?Div(?:ision)?\.?,? ?([12])\b|C1D([12])", str(na))
        if not div and "Class" not in str(na):
            div = re.search(r"Div(?:ision)?\.?,? ?([12])\b", str(na))
        if div:
            cert["division"] = f"Class I Division {div.group(1) or div.group(2)}"
        if "certification.northAmerica" in notes and "division" in cert:
            notes["certification.division"] = notes["certification.northAmerica"]

    epl_gas = best_epl(use(p, "EPL gas", "certification.epl", notes, ctx), "G")
    epl_dust = best_epl(use(p, "EPL dust", "certification.epl", notes, ctx), "D")
    cert["zonesGas"] = zones_from_epl(epl_gas, "gas")
    cert["zonesDust"] = zones_from_epl(epl_dust, "dust")
    if not cert["zonesGas"]:
        sys.exit(f"{pid}: no gas zones; it can't be listed")
    label, kind = protection(gi or ga or "")
    if label:
        cert["protection"] = label
    # Group and class come from the marking shown, so they always agree with it (e.g. Atexxo iPad mini: IIIA on the
    # Zone 21 marking, IIIB on the Zone 22 one); the cell is the fallback
    shown = {"gasGroup": gi or ga or "", "dustGroup": di or da or "", "tempClass": gi or ga or ""}
    for key, col, allowed in (("gasGroup", "Gas group", ("IIA", "IIB", "IIC")), ("dustGroup", "Dust group", ("IIIA", "IIIB", "IIIC")),
                              ("tempClass", "Temp class", ("T1", "T2", "T3", "T4", "T5", "T6"))):
        v = use(p, col, f"certification.{key}", notes, ctx)
        found = [a for a in allowed if re.search(rf"\b{a}\b", shown[key])] or [a for a in allowed if re.search(rf"\b{a}\b", str(v or ""))]
        if found:
            cert[key] = found[0]
    surface = use(p, "Max surface temp dust", "certification.maxSurfaceTemp", notes, ctx)
    if surface:
        cert["maxSurfaceTemp"] = re.sub(r"^T(\d+)\s*°C.*$", r"\1 °C", str(surface))
    epl = [e for e in (epl_gas, epl_dust) if e]
    if epl:
        cert["epl"] = " / ".join(epl)
    ta_text, ta_lo, ta_hi = temp_range(use(p, "Ex ambient Ta", "certification.ambient", notes, ctx))
    if ta_text:
        cert["ambient"] = ta_text
    if "0" in cert["zonesGas"]:
        cond = str(p.raw("Specific conditions / safety rules") or "") + " " + str(p.raw("Variant") or "")
        if re.search(r"leather case|carry case|only (?:in|inside)", cond, re.I):
            cert["zone0Note"] = "Zone 0 rated only inside the case supplied by the maker."

    # ----- Safe use
    safety = {}
    charging = use(p, "Charging in Ex area", "safety.charging", notes, ctx)
    if charging:
        if re.match(r"Yes\b", str(charging)):
            safety["chargingInZone"] = True
            safety["charging"] = clean(charging)
        else:
            safety["chargingInZone"] = False
            safety["charging"] = "Outside the hazardous area only"
    swap = use(p, "Battery swap in Ex area", "safety.batterySwap", notes, ctx)
    if swap and re.match(r"No\b", str(swap)):
        safety["batterySwap"] = "Outside the hazardous area only"
    elif swap and re.match(r"Yes\b", str(swap)):
        safety["batterySwap"] = clean(swap)
    conditions = use(p, "Specific conditions / safety rules", "safety.conditions", notes, ctx)
    if conditions:
        safety["conditions"] = minus(str(conditions))
    peripherals = use(p, "Ex-area restrictions on peripherals", "safety.peripherals", notes, ctx)
    if peripherals:
        safety["peripherals"] = minus(str(peripherals))
    accessories = use(p, "Certified accessories", "safety.certifiedAccessories", notes, ctx)
    if accessories:
        safety["certifiedAccessories"] = str(accessories)

    # ----- Operating system and device specs (from the host / base device)
    platform = str(base.raw("Platform") or "")
    os_name = "iPadOS" if "iPad" in platform else "Windows" if "Windows" in platform else "Android" if "Android" in platform else "Other"
    current_os = use(base, "Current OS", "os.version", notes)
    launch_os = use(base, "OS at launch", "os.atLaunch", notes)
    # Keep only the version itself ("Android 16 (Samsung lists an Android 16 manual)" -> "Android 16")
    os_text = lambda v: (re.match(r"^((?:Android|iPadOS|Windows)\b[^(;]*?)\s*(?:\(|;|$)", str(v)) or [None, str(v)])[1].strip() if v else None
    current_os, launch_os = (re.sub(r"\b(\d+)\.0\b", r"\1", x) if x else x for x in (os_text(current_os), os_text(launch_os)))
    version = current_os or launch_os or "Not published"
    if version in ("Android", "Windows"):
        version = f"{version} (version not published)"
    os_info = {"name": os_name, "version": str(version)}
    if launch_os and launch_os != version:
        os_info["atLaunch"] = str(launch_os)
    edition = use(base, "OS edition", "os.edition", notes)
    if edition:
        os_info["edition"] = str(edition)
    limit = use(base, "OS upgrade limit", "os.upgradeLimit", notes)
    if limit:
        os_info["upgradeLimit"] = str(limit)
    if os_name == "Windows" and re.search(r"Windows 11 (?:Pro|Professional|Home)", f"{edition} {version}"):
        os_info["modernLifecycle"] = True
        os_info["updates"] = "Windows 11: supported while kept on a current version (Microsoft Modern Lifecycle)"
    # Windows 11 Pro has no fixed end date, so the workbook's placeholder date isn't read (or noted) at all
    until = None if os_info.get("modernLifecycle") else iso_date(use(base, "Security updates until", "os.securityUntil", notes))
    if until:
        os_info["securityUntil"] = until
        if base.id in SECURITY_QUARTERS and quarter_label(until):
            os_info["securityUntilLabel"] = quarter_label(until)
        if base.id.startswith("B-APL-IPADPRO"):
            # Decision D2 (Rory, 2026-10-26), as the phones site does for iPhones: the workbook holds Apple's legal
            # minimum (first sale + 5 years); earlier iPad Pro models got 7-8 years, so publish that and move the date
            # on 2 years for sorting and filters.
            d = datetime.date.fromisoformat(until)
            d = d.replace(day=28) if (d.month, d.day) == (2, 29) else d
            os_info["securityUntil"] = d.replace(year=d.year + 2).isoformat()
            os_info["supportYears"] = IPAD_PRO_SUPPORT_YEARS
            os_info["updates"] = "Security updates for 7–8 years (based on previous iPad Pro support)"
        elif os_name == "iPadOS":
            os_info["updates"] = f"Security updates guaranteed until at least {month_year(until)} (Apple's minimum)"
        else:
            os_info["updates"] = f"Security updates {'ended in' if security_ended(os_info) else 'until'} {security_end(os_info)}"

    specs = {}
    size = first_number(base.raw("Display (in)"))
    dtype = base.get("Display type / glass")
    if size:
        specs["display"] = f"{size:g}-inch" + (f" {dtype}" if dtype else "")
    for key, col, fmt in [
        ("resolution", "Resolution", lambda v: str(v).replace(" x ", " × ")),
        ("brightness", "Peak brightness (nits)", lambda v: f"{v} nits" if re.fullmatch(r"\d+", str(v)) else str(v)),
        ("outdoor", "Outdoor readability", str),
        ("chip", "Chipset", str),
        ("cpuArch", "CPU architecture", str),
        ("ram", "RAM (GB)", lambda v: f"{v} GB" if re.fullmatch(r"[\d.]+", str(v)) else str(v)),
        ("camera", "Rear camera", str),
        ("frontCamera", "Front camera", str),
        ("video", "Video", str),
        ("battery", "Battery (mAh)", lambda v: v if isinstance(v, str) else f"{v} mAh"),
        ("batteryWh", "Battery (Wh)", lambda v: v if isinstance(v, str) else f"{v} Wh"),
        ("hotSwap", "Hot-swap / bridge battery", str),
        ("batteryLife", "Battery life (maker figure)", str),
        ("ports", "Ports and expansion", str),
        ("docking", "Docking", str),
        ("network", "Network", str),
        ("wifi", "Wi-Fi", str),
        ("bluetooth", "Bluetooth", lambda v: f"Bluetooth {v}" if re.match(r"^[\d.]+$", str(v)) else str(v)),
        ("nfc", "NFC", str),
        ("gnss", "GNSS", str),
        ("sim", "SIM", str),
        ("keys", "Dedicated keys", str),
        ("biometrics", "Biometrics", str),
        ("depthSensing", "Depth sensing", str),
        ("touch", "Glove / wet touch", str),
        ("pen", "Pen / digitiser", str),
        ("keyboard", "Keyboard / 2-in-1", str),
        ("apps", "Preinstalled Ex-relevant apps", str),
    ]:
        v = use(base, col, f"specs.{key}", notes)
        if v is not None and not (key == "depthSensing" and str(v).startswith(("None", "No"))):
            specs[key] = fmt(v)
    storage = base.get("Storage (GB)")
    if storage:
        s = re.sub(r"\b(\d+)\b(?!\s*(?:GB|TB))", r"\1 GB", str(storage))
        expand = base.get("Expandable")
        specs["storage"] = s + (f", {expand}" if expand and str(expand).lower().startswith("microsd") else "")
    removable = base.get("Removable battery")
    if removable:
        specs["replaceableBattery"] = str(removable).startswith("Yes")
    if specs.get("touch"):
        specs["gloveTouch"] = bool(re.search(r"glove", specs["touch"], re.I))
    if specs.get("keys"):
        specs["pttKey"] = bool(re.search(r"\bPTT\b|push-to-talk", specs["keys"], re.I))
    specs["connectivity"] = [x for x in (
        # "4G LTE (no 5G)" must not read as 5G
        "5G" if re.search(r"\b5G\b", re.sub(r"\bno 5G\b", "", specs.get("network", ""))) else
        "4G" if re.search(r"4G|LTE", specs.get("network", "")) else None,
        (re.match(r"Wi-Fi \w+", specs.get("wifi", "")) or [None])[0],
        "NFC" if specs.get("nfc", "").startswith("Yes") else None,
        "eSIM" if "eSIM" in specs.get("sim", "") else None,
    ) if x]

    # Product-level physical facts
    weight = number(use(p, "In-use weight (g)", "specs.weightG", notes, ctx))
    if weight:
        specs["weightG"] = weight
        if ttype in ("ex-case", "conversion"):
            specs["weightNote"] = "Including the tablet"
    dims = use(p, "Dimensions (mm)", "specs.dimensionsMm", notes, ctx)
    if dims:
        specs["dimensionsMm"] = str(dims).replace(" x ", " × ")
    ip = use(p, "IP rating", "specs.ipRating", notes, ctx)
    if ip:
        specs["ipRating"] = str(ip)
    exip = use(p, "Ex IP (cert)", "specs.exIp", notes, ctx)
    if exip:
        specs["exIp"] = str(exip)
    drop = use(p, "Drop / MIL-STD", "specs.ruggedness", notes, ctx)
    if drop:
        specs["ruggedness"] = str(drop)
    op_text, _, _ = temp_range(use(p, "Operating temp (°C)", "specs.operatingTemp", notes, ctx))
    if op_text:
        specs["operatingTemp"] = op_text
    camera_free = use(p, "Camera-free variant", "specs.cameraFree", notes, ctx)
    if camera_free:
        if str(camera_free).startswith("Yes"):
            specs["cameraFree"] = True
        elif re.match(r"No\b", str(camera_free)):
            specs["cameraFree"] = False
    parts = use(p, "Part numbers / SKUs", "specs.partNumbers", notes, ctx)
    if parts:
        specs["partNumbers"] = str(parts)

    # ----- Hosts, fitting, removability
    hosts = [host_name(bases[b].raw("Name")) for b in base_ids if b in bases]
    fitting = use(p, "Fitting", "fitting", notes, ctx)
    removable_text = use(p, "Tablet removable and reusable", "removable", notes, ctx)

    # ----- What works through the case, and field capabilities
    through_rows = []
    t = through.get(pid)
    if t and ttype in ("ex-case", "conversion"):
        for col, label in THROUGH_LABELS.items():
            v = use(t, col, f"throughCase.{label}", notes)
            if v is None or str(v).startswith("N/A"):
                continue
            through_rows.append({"label": label, "value": tidy(v)})
    cap_rows = []
    c = caps.get(pid)
    if c:
        for col, label in CAPABILITY_LABELS.items():
            v = use(c, col, f"capabilities.{label}", notes)
            if v is None or str(v).startswith("Not listed"):
                continue
            cap_rows.append({"label": label, "value": tidy(v)})
    cap = lambda label: next((r["value"] for r in cap_rows if r["label"] == label), "")

    # ----- Versions folded into this page
    variants = list(extra_variants)
    for child in variants_by_parent.get(pid, []):
        vname, text = VARIANT_TEXT[child.id]
        variant = {"name": vname, "description": text}
        mining = child.get("Marking - mining")
        if mining and child.id in MINING_VARIANTS:
            variant["marking"] = minus(str(mining))
        numbers = [split_scope(str(n))[0] for n in (child.get("ATEX certificate"), child.get("IECEx certificate"))
                   if n and re.search(r"ATEX|IECEx", str(n))]
        if numbers:
            variant["certificates"] = "; ".join(numbers)
        variants.append(variant)

    # ----- Features the site can't derive from the specs alone
    features = []
    if cap("NFC reading").startswith("Built in") and not HEDGED.search(cap("NFC reading")):
        features.append("nfc")
    if specs.get("cameraFree"):
        features.append("camera-free")
    if "LiDAR" in specs.get("depthSensing", "") and (ttype not in ("ex-case", "conversion") or
                                                     any(r["label"] == "LiDAR" and r["value"].startswith("Yes") for r in through_rows)):
        features.append("lidar")
    if re.match(r"Built in|Via accessory", cap("Thermal imaging")):
        features.append("thermal")
    barcode = cap("Barcode scanning")
    if (barcode.startswith("Built in") and re.search(r"imager|scan engine|scanner|reader", barcode, re.I)) or (
            barcode.startswith("Via accessory") and re.search(r"Ex-rated\? yes", barcode)):
        features.append("barcode")
    if safety.get("chargingInZone"):
        features.append("charge-in-zone")
    if (os_info.get("securityUntil") or "") >= "2030" or os_info.get("modernLifecycle"):
        features.append("updates-2030")
    if any(ch.id in MINING_VARIANTS for ch in variants_by_parent.get(pid, [])):
        features.append("mining-version")

    # ----- Sources
    ids = source_ids(p.raw("Sources"), base.raw("Sources"), t.raw("Sources") if t else "", c.raw("Sources") if c else "")
    srcs = []
    for sid in ids:
        s = sources.get(sid)
        url = s and s.raw("URL")
        if not url or not str(url).startswith("http"):
            continue
        rank = SOURCE_ORDER.index(s.raw("Type")) if s.raw("Type") in SOURCE_ORDER else len(SOURCE_ORDER)
        srcs.append((rank, {"label": f"{s.raw('Publisher')}: {s.raw('Document')}", "url": str(url)}))
    srcs = [s for _, s in sorted(srcs, key=lambda x: x[0])][:25]
    if not srcs and website:
        srcs = [{"label": f"{maker} website", "url": website}]

    manufacturer = {"name": maker, "country": country}
    if website:
        manufacturer["website"] = website
    tab = {
        "productId": pid,
        "name": name,
        "brand": brand_name,
        "brandSlug": brand_slug,
        "manufacturer": manufacturer,
        "type": ttype,
    }
    if size:
        tab["screenIn"] = int(size) if float(size).is_integer() else size
    if ttype in ("ex-case", "conversion"):
        tab["hostDevice"] = hosts[0]
        tab["otherHosts"] = hosts[1:]
        if fitting:
            tab["fitting"] = ("unit" if re.search(r"unit", str(fitting), re.I) else
                              "self-install" if re.search(r"self", str(fitting), re.I) else "factory-fitted")
        if removable_text:
            tab["removable"] = clean(removable_text)
    tab["status"] = status
    if status_note:
        tab["statusNote"] = status_note
    if successor:
        tab["successor"] = PRODUCTS[successor][0]
    launch = use(p, "Launch", "launchYear", notes, ctx)
    if launch and re.match(r"\d{4}", str(launch)):
        tab["launchYear"] = int(str(launch)[:4])
    if ttype in ("ex-case", "conversion"):
        # The base tablet's user-replaceable or hot-swap battery doesn't survive a sealed Ex housing unless the maker
        # says so (the housing's own rule is safety.batterySwap), so the host's battery claims aren't published
        if specs.get("replaceableBattery"):
            del specs["replaceableBattery"]
        if specs.get("hotSwap", "").lower().startswith("yes"):
            del specs["hotSwap"]
    tab.update({
        "summary": "",
        **({k: existing[k] for k in ("image", "imageAlt") if k in existing}),
        "certification": cert,
        "safety": safety,
        "os": os_info,
        "specs": specs,
        "throughCase": through_rows,
        "capabilities": cap_rows,
        "variants": variants,
        "unconfirmed": notes,
        "features": features,
        "strengths": [],
        "limitations": [],
        "bestFor": [],
        "alternatives": [],
        "sponsor": brand_slug in SPONSOR_BRANDS,
        "lastVerified": iso_date(p.raw("Date verified")) or TODAY,
        "sources": srcs,
    })
    unverified = "certification.atex.marking" in notes and not (cert.get("atex", {}).get("number") or cert.get("iecex", {}).get("number"))
    tab["_ctx"] = {"kind": kind, "ta": (ta_lo, ta_hi), "marking_unverified": unverified,
                   "apple_min_ended": os_name == "iPadOS" and bool(until) and until < TODAY,
                   "arm_windows": os_name == "Windows" and "Arm" in str(specs.get("cpuArch", ""))}
    return tab


# ---------------------------------------------------------------------------------------------------------------
# Copy: the same rules for every brand
# ---------------------------------------------------------------------------------------------------------------

def schemes(t):
    c = t["certification"]
    return join_and([c.get("atex") and "ATEX", c.get("iecex") and "IECEx", c.get("ukex") and "UKEX"])


def scheme_phrase(t):
    """" under ATEX and IECEx" for the summary. An ATEX rating with no published certificate isn't called certified when
    another scheme's certificate is, so the sentence doesn't claim more than the documents show."""
    c = t["certification"]
    unpublished = c.get("atex", {}).get("kind") == "not-published"
    shown = join_and([c.get("atex") and not unpublished and "ATEX", c.get("iecex") and "IECEx", c.get("ukex") and "UKEX"])
    if unpublished and shown:
        return f" under {shown}. {t['brand']} also states ATEX certification, but the certificate isn't published"
    return f" under {schemes(t)}" if schemes(t) else ""


def protection_phrase(t):
    """Adjective for the product's protection. Only Ex i products are called intrinsically safe."""
    kind = t["_ctx"]["kind"]
    m = re.search(r"\(Ex ([^)]+)\)", t["certification"].get("protection", ""))
    concept = m.group(1) if m else ""
    if t["_ctx"]["marking_unverified"]:
        # Only a marking backed by a certificate earns a protection term, as on the phones site
        return f"Ex {concept}-rated" if concept else "Ex-rated"
    return {"intrinsic": f"intrinsically safe (Ex {concept})", "flameproof": "flameproof (Ex d)",
            "restricted": "restricted-breathing (Ex nR)", "increased": f"increased-safety (Ex {concept})",
            "combined": f"Ex {concept}"}.get(kind, "Ex-certified")


def summary(t, names):
    c, o = t["certification"], t["os"]
    zones = zone_label(c["zonesGas"], c["zonesDust"])
    verb = "rated for" if t["_ctx"]["marking_unverified"] else "certified for"
    under = scheme_phrase(t)
    adj = protection_phrase(t)
    size = f"{t['screenIn']:g}-inch " if t.get("screenIn") else ""
    os_v = o["version"]
    runs = f"It runs {os_v}." if os_v != "Not published" else "Its operating system version isn't published."
    name = t["name"] if t["name"].startswith("The ") else f"The {t['name']}"
    if t["type"] == "dedicated":
        text = f"{name} is {article(adj)} {adj} {size}{o['name']} tablet from {t['brand']}, {verb} {zones}{under}. {runs}"
    elif t["type"] == "rugged-variant":
        text = (f"{name} is the hazardous-area version of a {t['brand']} rugged {size}tablet, {verb} {zones}{under}. "
                f"It is {article(adj)} {adj} design" + (f" running {os_v}." if os_v != "Not published" else "; its operating system version isn't published."))
    elif universal(t):
        text = (f"{name} is a universal {adj} case for 10- to 11-inch tablets ({host_makes(t)}), {verb} {zones}{under}. "
                f"It takes a standard tablet from the supported list, so it works with the same apps and device management.")
    elif t["type"] == "ex-case":
        text = (f"{name} is {article(adj)} {adj} case for the {t['hostDevice']}, {verb} {zones}{under}. "
                f"The tablet inside is a standard {t['hostDevice']} running {os_v}, so it works with the same apps and device management.")
    else:
        text = (f"{name} is a standard {t['hostDevice']} converted by {t['brand']} for hazardous areas, {verb} "
                f"{zones}{under}. It is {article(adj)} {adj} design, and the tablet runs {os_v}.")
    if t["status"] == "discontinued":
        succ = t.get("successor")
        text += f" It has been discontinued; the {names[succ]} replaces it." if succ else " It has been discontinued."
    return text


def universal(t):
    """A case certified for tablets from more than one maker (Armadex Ex UniCase)"""
    return t["type"] == "ex-case" and len({h.split()[0] for h in [t["hostDevice"], *t["otherHosts"]]}) > 1


def host_makes(t):
    return join_and(sorted({"iPad" if h.startswith("iPad") else "Galaxy Tab" if "Galaxy" in h else "Surface Go" if "Surface" in h else h
                            for h in [t["hostDevice"], *t["otherHosts"]]}), )


def strengths(t, lighter):
    c, s, o, u = t["certification"], t["specs"], t["os"], t["unconfirmed"]
    out = []
    if "0" in c["zonesGas"]:
        out.append("Zone 0 rated when used in the supplied case" if c.get("zone0Note") else "Certified for Zone 0, the highest-risk gas zone")
    if t["safety"].get("chargingInZone"):
        out.append("May be charged inside the hazardous area, within the certificate's limits")
    if re.search(r"data[^.]{0,30} in the (?:hazardous area|zone)", t["safety"].get("peripherals", ""), re.I):
        out.append("USB-C can be used for data in the zone")
    if c.get("atex") and c.get("iecex") and c["atex"].get("kind") != "not-published":
        out.append(f"{schemes(t)} certified" + (", with published certificate numbers" if c["atex"].get("number") and c["iecex"].get("number") else ""))
    if c.get("division") and "certification.division" not in u:
        out.append(f"Also certified for {c['division']} in North America")
    if c.get("tempClass") == "T6" and "certification.tempClass" not in u:
        out.append("T6 temperature class, the strictest rating")
    if "mining-version" in t["features"]:
        out.append("A mining version (Group I) is also available")
    if "barcode" in t["features"]:
        bc = next((r["value"] for r in t["capabilities"] if r["label"] == "Barcode scanning"), "")
        out.append("Ex-rated barcode scanner available as an accessory" if bc.startswith("Via accessory")
                   else "Optional built-in barcode imager" if "optional" in bc.lower() else "Built-in barcode imager")
    if "5G" in s.get("connectivity", []):
        optional = re.search(r"optional|Wi-Fi only|\bor\b", s.get("network", ""), re.I)
        out.append("5G on cellular versions" if optional else "5G connectivity")
    if t["status"] == "current" and o.get("securityUntil", "") >= "2030" and not o.get("modernLifecycle") and not universal(t):
        out.append(o["updates"])
    if t["type"] in ("dedicated", "rugged-variant") and s.get("replaceableBattery"):
        out.append("User-replaceable battery")
    if s.get("hotSwap", "").startswith("Yes"):
        out.append("Hot-swappable battery for long shifts")
    if s.get("cameraFree"):
        out.append("Camera-free version available for sites that restrict cameras")
    if "lidar" in t["features"]:
        out.append("LiDAR scanner usable through the case")
    if universal(t):
        out.append(f"Takes standard 10- to 11-inch tablets ({host_makes(t).replace(' and ', ' or ')}), so you keep your usual apps")
    elif t["type"] in ("ex-case", "conversion"):
        mdm = "Apple device management" if o["name"] == "iPadOS" else "standard device management"
        out.append(f"A standard {t['hostDevice']} inside, with the same apps and {mdm}")
    if lighter:
        out.append(lighter)
    hi = t["_ctx"]["ta"][1]
    if hi is not None and hi >= 60:
        out.append("Certified for ambient temperatures up to +60 °C")
    if not out:
        out.append(f"Certified for {zone_label(c['zonesGas'], c['zonesDust'])}")
    return out[:5]


def limitations(t, names):
    c, s, o, u = t["certification"], t["specs"], t["os"], t["unconfirmed"]
    ctx = t["_ctx"]
    out = []
    if t["status"] == "discontinued":
        succ = t.get("successor")
        out.append(f"Discontinued; the {names[succ]} replaces it" if succ else "No longer in the maker's current range")
    if c["zonesGas"] == ["2"]:
        out.append("Zone 2 only, so not for Zone 1 areas")
    if c.get("gasGroup") == "IIB":
        out.append("Gas group IIB: not rated for hydrogen or acetylene")
    if ctx["marking_unverified"]:
        out.append("Ex marking not yet confirmed against a published certificate")
    tc = c.get("tempClass")
    if tc in ("T3", "T4", "T5") and "certification.tempClass" not in u:
        limit = {"T3": 200, "T4": 135, "T5": 100}[tc]
        out.append(f"{tc} temperature class: fine for most gases, but not those that ignite below {limit} °C")
    if not c["zonesDust"]:
        out.append("No dust zone rating")
    elif c.get("dustGroup") == "IIIB":
        out.append("Dust group IIIB: combustible flyings and non-conductive dust only, not conductive dust")
    elif c.get("dustGroup") == "IIIA":
        out.append("Dust group IIIA: combustible flyings only, not dusts")
    if c.get("atex", {}).get("kind") == "unit-assessment":
        out.append("Certified by ATEX unit assessment; no certificate number is published")
    if c.get("atex") and not c.get("iecex"):
        out.append("ATEX only: no IECEx certificate")
    if ctx["arm_windows"]:
        out.append("Windows on Arm: x86/x64 apps run through emulation and need Arm64 drivers")
    if ctx["apple_min_ended"]:
        out.append("Apple's minimum update period for this iPad has ended: for existing fleets only")
    lo, hi = ctx["ta"]
    if hi is not None and hi < 50:
        out.append(f"Ambient rating tops out at +{hi} °C")
    if lo is not None and lo > -20:
        out.append(f"Not certified below {str(lo).replace('-', '−')} °C ambient")
    m = re.match(r"Android (\d+)", o.get("upgradeLimit") or o.get("version") or "")
    if security_ended(o) and o["name"] != "iPadOS":
        out.append(f"Security updates ended in {security_end(o)}")
    elif t["status"] == "current" and o.get("securityUntil") and not o.get("modernLifecycle") and o["securityUntil"] < SOON and o["name"] != "iPadOS":
        out.append(f"Security updates end in {security_end(o)}")
    elif t["status"] == "current" and o["name"] == "Android" and not o.get("securityUntil"):
        out.append("No published end date for security updates" + (f" (runs Android {m.group(1)})" if m and int(m.group(1)) <= 11 else ""))
    if re.search(r"Windows 10", o.get("edition", "") + o["version"]) and t["status"] == "current" and not o.get("securityUntil"):
        out.append("Ships with Windows 10: check which edition and its support end date")
    if t["type"] in ("dedicated", "rugged-variant") and s.get("replaceableBattery") is False:
        out.append("Battery isn't user-replaceable")
    if not s.get("weightG"):
        out.append("In-use weight not published")
    if "0" not in c["zonesGas"] and "1" in c["zonesGas"]:
        out.append("Not certified for Zone 0")
    return out[:5]


def top_zone(t):
    zs = [z for z in t["certification"]["zonesGas"] if z != "0"] or t["certification"]["zonesGas"]
    return sorted(zs)[0]


def lighter_note(t, tabs):
    w, size = t["specs"].get("weightG"), t.get("screenIn")
    if not w or not size:
        return None
    group = [q["specs"]["weightG"] for q in tabs if q["status"] == "current" and q["specs"].get("weightG") and q.get("screenIn")
             and abs(q["screenIn"] - size) <= 1.5 and top_zone(q) == top_zone(t)]
    if len(group) < 4:
        return None
    group.sort()
    if w <= group[max(0, len(group) // 3 - 1)]:
        return f"Among the lightest Zone {top_zone(t)} tablets of its size we list, at {w} g in use"
    return None


def alternatives(t, tabs):
    family = lambda q: re.sub(r"\s+(DZ\d|EX\d|\.\d|G\d.*|Zone \d).*$", "", q["name"])
    cands = []
    for q in tabs:
        if q is t or q["status"] != "current" or (q["brandSlug"] == t["brandSlug"] and family(q) == family(t)):
            continue
        score = 10 if top_zone(t) in q["certification"]["zonesGas"] else 0
        score += 2 if top_zone(q) == top_zone(t) else 0
        score += 4 if q["os"]["name"] == t["os"]["name"] else 0
        score += 2 if q["type"] == t["type"] else 0
        score += 2 if q["brandSlug"] != t["brandSlug"] else -8
        if q.get("screenIn") and t.get("screenIn"):
            score -= abs(q["screenIn"] - t["screenIn"]) / 2
        cands.append((score, q["name"], q))
    cands.sort(key=lambda x: (-x[0], x[1]))
    picked, brands = [], set()
    if t.get("successor"):
        picked.append(t["successor"])
        brands.add(t["brandSlug"])
    for _, _, q in cands:
        if len(picked) >= 3:
            break
        if q["brandSlug"] in brands or q["slug"] in picked:
            continue
        picked.append(q["slug"])
        brands.add(q["brandSlug"])
    return picked


def best_for(t):
    """Best for pages whose pool the tablet belongs to, linked from its page once the Best for page exists."""
    if t["status"] != "current":
        return []
    out = []
    if top_zone(t) == "1":
        out.append("zone-1-tablets")
    if t["certification"]["zonesGas"] == ["2"]:
        out.append("zone-2-tablets")
    if t["certification"].get("division") and "certification.division" not in t["unconfirmed"]:
        out.append("division-1-tablets")
    if "mining-version" in t["features"]:
        out.append("mining")
    out.append({"iPadOS": "ipad", "Windows": "windows-tablets", "Android": "android-tablets"}.get(t["os"]["name"], "other"))
    return out


# ---------------------------------------------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--report", help="write a Markdown report of every public note and skipped cell; don't write JSON")
    args = ap.parse_args()
    products, bases, through, caps, sources = load(args.workbook)

    variants_by_parent, skipped = {}, []
    for pid, p in products.items():
        depth, status = str(p.raw("Page depth") or ""), str(p.raw("Status") or "")
        if depth.startswith("Variant of"):
            parent = depth.split("Variant of", 1)[1].strip()
            if pid not in VARIANT_TEXT:
                sys.exit(f"{pid} is a variant: add it to VARIANT_TEXT")
            variants_by_parent.setdefault(parent, []).append(p)
        elif status.startswith("Announced"):
            skipped.append(pid)
        elif pid not in PRODUCTS:
            sys.exit(f"{pid} isn't mapped: add it to PRODUCTS (slug and name), make it a variant, or mark it Announced")
        elif p.raw("Brand") not in BRANDS:
            sys.exit(f"{pid}: brand {p.raw('Brand')!r} isn't in BRANDS")

    tabs = []
    for pid, (slug, _) in PRODUCTS.items():
        if pid not in products:
            sys.exit(f"{pid} is in PRODUCTS but not in the workbook")
        path = OUT / f"{slug}.json"
        existing = json.loads(path.read_text()) if path.exists() else {}
        t = build(products[pid], bases, through, caps, sources, variants_by_parent, existing)
        t["slug"] = slug
        tabs.append(t)

    if args.report:
        write_report(Path(args.report), tabs)
        return
    if MISSING:
        lines = "\n".join(f"  {s} | {rid} | {col} | {str(v)[:60]}" for s, rid, col, v, _ in MISSING)
        sys.exit(f"Flagged values without a public note (add a rule to NOTE_RULES, FIELD_NOTES or STATUS):\n{lines}")

    names = {t["slug"]: t["name"] for t in tabs}
    for t in tabs:
        t["summary"] = summary(t, names)
        t["strengths"] = strengths(t, lighter_note(t, tabs))
        t["limitations"] = limitations(t, names)
        t["alternatives"] = alternatives(t, tabs)
        t["bestFor"] = best_for(t)

    OUT.mkdir(parents=True, exist_ok=True)
    written = set()
    for t in tabs:
        slug = t.pop("slug")
        t.pop("_ctx")
        (OUT / f"{slug}.json").write_text(json.dumps(t, ensure_ascii=False, indent=2) + "\n")
        written.add(f"{slug}.json")
    removed = [f.name for f in OUT.glob("*.json") if f.name not in written]
    for name in removed:
        (OUT / name).unlink()
    print(f"Wrote {len(tabs)} tablets to {OUT.relative_to(ROOT)} ({sum(len(t['unconfirmed']) for t in tabs)} values with a public note)")
    print(f"Listed as variants: {', '.join(sorted(c.id for v in variants_by_parent.values() for c in v))}")
    print(f"Left out until certified: {', '.join(skipped) or 'none'}")
    if removed:
        print(f"Removed: {', '.join(removed)}")


def write_report(path, tabs):
    by_note = Counter(n for *_, n in NOTES_USED)
    lines = [f"# Public notes: dry run ({TODAY})", "",
             f"{len(tabs)} tablets; {len(NOTES_USED)} flagged values would be published with a note; "
             f"{len(SKIPPED)} flagged values are not published; {len(MISSING)} flagged values have no note yet.", "",
             "## Notes in use (count, text)", ""]
    lines += [f"- **{n}×** {text}" for text, n in by_note.most_common()]
    lines += ["", "## Not published (skip rules)", ""]
    lines += [f"- {rid} | {col}: {reason}" for _, rid, col, reason in SKIPPED]
    lines += ["", "## Still without a note", ""]
    lines += [f"- {s} | {rid} | {col} | `{str(v)[:60]}` | researcher: {str(com)[:160]}" for s, rid, col, v, com in MISSING]
    path.write_text("\n".join(lines) + "\n")
    print(f"Report written to {path}: {len(NOTES_USED)} notes, {len(SKIPPED)} skipped, {len(MISSING)} missing")


if __name__ == "__main__":
    main()
